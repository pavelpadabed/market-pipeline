from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from market_pipeline.presentation.models import (
    LatestOfferReportRow,
)


class ExcelLatestOffersExporter:
    def export(
        self,
        report_rows: tuple[LatestOfferReportRow, ...],
        path: Path,
    ) -> None:
        headers = (
            "Product",
            "Store",
            "Price",
            "Currency",
            "Observed At (Europe/Warsaw)",
        )

        column_widths = {
            "A": 36,
            "B": 20,
            "C": 12,
            "D": 12,
            "E": 24,
        }

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(headers)
        worksheet.freeze_panes = "A2"
        horizontal_center_alignment = Alignment(horizontal="center")
        for column_letter, width in column_widths.items():
            worksheet.column_dimensions[column_letter].width = width

        header_font = Font(bold=True, color="FFFFFFFF")
        header_fill = PatternFill(fill_type="solid", fgColor="FF1F4E78")

        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = horizontal_center_alignment

        warsaw_timezone = ZoneInfo("Europe/Warsaw")

        if not report_rows:
            message = "No offers were found in the latest run."
            worksheet["A2"].value = message

        for report_row in report_rows:
            excel_observed_at = (
                report_row.observed_at.astimezone(warsaw_timezone).replace(
                    tzinfo=None,
                )
            )
            report_values = (
                report_row.product_title,
                report_row.store_name,
                report_row.price,
                report_row.currency,
                excel_observed_at,
            )

            worksheet.append(report_values)
            excel_row = worksheet.max_row

            store_cell = worksheet.cell(
                row=excel_row,
                column=2,
            )
            store_cell.alignment = horizontal_center_alignment

            price_cell = worksheet.cell(
                row=excel_row,
                column=3,
            )
            price_cell.number_format = "0.00"

            currency_cell = worksheet.cell(
                row=excel_row,
                column=4,
            )
            currency_cell.alignment = horizontal_center_alignment

            observed_at_cell = worksheet.cell(
                row=excel_row,
                column=5,
            )
            observed_at_cell.number_format = "yyyy-mm-dd hh:mm:ss"
            observed_at_cell.alignment = horizontal_center_alignment

        worksheet.auto_filter.ref = f"A1:E{1 + len(report_rows)}"

        workbook.save(path)
