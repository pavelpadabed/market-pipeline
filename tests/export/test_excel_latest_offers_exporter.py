from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from market_pipeline.export.excel_latest_offers_exporter import (
    ExcelLatestOffersExporter,
)
from market_pipeline.presentation.models import LatestOfferReportRow


def _make_report_row(**overrides: object) -> LatestOfferReportRow:
    data = {
        "product_title": "Batman",
        "store_name": "Steam",
        "price": Decimal("10.13"),
        "currency": "USD",
        "observed_at": datetime(
            2026, 9, 22, 6, 52, tzinfo=UTC,
        ),
    }
    data.update(overrides)

    return LatestOfferReportRow(**data)


def test_export_writes_report_row_to_excel_file(tmp_path: Path) -> None:
    workbook_path = tmp_path / "latest_offers.xlsx"

    report_row = LatestOfferReportRow(
        product_title="Batman",
        store_name="Steam",
        price=Decimal("10.13"),
        currency="USD",
        observed_at=datetime(
            2026, 9, 19, 15, 48, tzinfo=UTC,
        ),
    )
    expected_observed_at = datetime(2026, 9, 19, 17, 48)
    expected_values = (
        (
            "Product",
            "Store",
            "Price",
            "Currency",
            "Observed At (Europe/Warsaw)",
        ),
        ("Batman", "Steam", 10.13, "USD", expected_observed_at),
    )
    exporter = ExcelLatestOffersExporter()

    exporter.export((report_row,), workbook_path)

    workbook = load_workbook(workbook_path)
    actual_values = tuple(workbook.active.values)
    workbook.close()

    assert actual_values == expected_values


def test_export_styles_header_row(tmp_path: Path) -> None:
    workbook_path = tmp_path / "latest_offers.xlsx"

    report_rows = ()
    exporter = ExcelLatestOffersExporter()

    exporter.export(report_rows, workbook_path)

    workbook = load_workbook(workbook_path)

    try:
        for row in workbook.active["A1:E1"]:
            for cell in row:
                assert cell.font.bold is True
                assert cell.font.color.rgb == "FFFFFFFF"
                assert cell.fill.fill_type == "solid"
                assert cell.fill.fgColor.rgb == "FF1F4E78"
    finally:
        workbook.close()


def test_export_formats_price_with_two_decimal_places(tmp_path: Path) -> None:
    workbook_path = tmp_path / "latest_offers.xlsx"

    report_row = _make_report_row()

    exporter = ExcelLatestOffersExporter()

    exporter.export((report_row,), workbook_path)

    workbook = load_workbook(workbook_path)

    actual_price_number_format = workbook.active["C2"].number_format

    workbook.close()

    assert actual_price_number_format == "0.00"


def test_export_formats_observed_at_as_datetime(tmp_path: Path) -> None:
    workbook_path = tmp_path / "latest_offers.xlsx"

    report_row = _make_report_row()

    exporter = ExcelLatestOffersExporter()

    exporter.export((report_row,), workbook_path)

    workbook = load_workbook(workbook_path)

    actual_datetime_number_format = workbook.active["E2"].number_format

    workbook.close()

    assert actual_datetime_number_format == "yyyy-mm-dd hh:mm:ss"


def test_export_freezes_header_row(tmp_path: Path) -> None:
    workbook_path = tmp_path / "latest_offers.xlsx"

    report_row = _make_report_row()

    exporter = ExcelLatestOffersExporter()

    exporter.export((report_row,), workbook_path)

    workbook = load_workbook(workbook_path)

    actual_freeze_panes = workbook.active.freeze_panes

    workbook.close()

    assert actual_freeze_panes == "A2"


def test_export_enables_filter_for_report_range(tmp_path: Path) -> None:
    workbook_path = tmp_path / "latest_offers.xlsx"

    steam_row = _make_report_row()

    gamers_gate_row = _make_report_row(
        store_name="GamersGate",
        price=Decimal("7.34"),
    )

    exporter = ExcelLatestOffersExporter()

    exporter.export(
        (
            steam_row,
            gamers_gate_row,
        ),
        workbook_path,
    )

    workbook = load_workbook(workbook_path)

    actual_filter_range = workbook.active.auto_filter.ref

    workbook.close()

    assert actual_filter_range == "A1:E3"


def test_export_sets_report_column_widths(tmp_path: Path) -> None:
    workbook_path = tmp_path / "latest_offers.xlsx"

    expected_widths = {
        "A": 36,
        "B": 20,
        "C": 12,
        "D": 12,
        "E": 24,
    }

    report_rows = ()

    exporter = ExcelLatestOffersExporter()

    exporter.export(report_rows, workbook_path)

    workbook = load_workbook(workbook_path)

    actual_widths = {
        column_letter: workbook.active.column_dimensions[column_letter].width
        for column_letter in expected_widths
    }

    workbook.close()

    assert actual_widths == expected_widths


def test_export_writes_message_for_empty_report(tmp_path: Path) -> None:
    workbook_path = tmp_path / "latest_offers.xlsx"

    report_rows = ()

    exporter = ExcelLatestOffersExporter()

    exporter.export(report_rows, workbook_path)

    workbook = load_workbook(workbook_path)

    actual_message = workbook.active["A2"].value
    empty_cells, = workbook.active["B2:E2"]
    actual_empty_values = tuple(cell.value for cell in empty_cells)
    actual_filter_range = workbook.active.auto_filter.ref

    workbook.close()

    assert actual_message == "No offers were found in the latest run."
    assert actual_empty_values == (None, None, None, None)
    assert actual_filter_range == "A1:E1"


def test_exporter_centers_headers_and_selected_report_values(tmp_path: Path) -> None:
    workbook_path = tmp_path / "latest_offers.xlsx"

    report_row = _make_report_row()

    exporter = ExcelLatestOffersExporter()

    exporter.export((report_row,), workbook_path)

    workbook = load_workbook(workbook_path)

    actual_header_alignments = tuple(
        cell.alignment.horizontal
        for cell in workbook.active[1]
    )
    actual_store_alignment = workbook.active["B2"].alignment.horizontal
    actual_currency_alignment = workbook.active["D2"].alignment.horizontal
    actual_observed_at_alignment = workbook.active["E2"].alignment.horizontal

    workbook.close()

    actual_report_alignments = (
        actual_store_alignment,
        actual_currency_alignment,
        actual_observed_at_alignment,
    )

    assert all(
        value == "center"
        for value in actual_header_alignments
    )

    assert all(
        value == "center"
        for value in actual_report_alignments
    )
